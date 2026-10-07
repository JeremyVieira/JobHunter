from __future__ import annotations

import re
import warnings
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from hashlib import sha256
from html import escape
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate

try:
    from jobhunter.config.settings import OUTPUT_DIR, USER_PROFILE
except ModuleNotFoundError:  # pragma: no cover
    from config.settings import OUTPUT_DIR, USER_PROFILE

SECTION_HEADERS = {
    "PROFESSIONAL SUMMARY",
    "EXPERIENCE",
    "EDUCATION",
    "PROJECTS",
    "TECHNICAL SKILLS",
    "LANGUAGES",
    "PROFIL PROFESSIONNEL",
    "EXPÉRIENCE",
    "FORMATION",
    "COMPÉTENCES TECHNIQUES",
    "LANGUES",
}
# Sections whose body text is prose (a paragraph to read), not a job-title/company
# line. The parser needs this distinction so the summary doesn't get bolded like
# an experience entry.
SUMMARY_SECTION_HEADERS = {"PROFESSIONAL SUMMARY", "PROFIL PROFESSIONNEL"}
# Short factual lines (not prose, not a "Company - Title" heading) that should
# render as plain text rather than a bold entry heading.
PLAIN_TEXT_SECTION_HEADERS = SUMMARY_SECTION_HEADERS | {"LANGUAGES", "LANGUES"}
SKILLS_SECTION_HEADERS = {"TECHNICAL SKILLS", "COMPÉTENCES TECHNIQUES"}
# A date-range line ("April 2024 - Current", "Septembre 2022 - Juin 2026") always
# contains a year; a "Company, City - Title" entry line never does. Using the
# dash style itself to tell them apart is fragile -- the English master resume
# happens to use an em-dash ("—") for entries, but the French one uses a plain
# hyphen ("Montréal - Analyste..."), which collided with the old " - " check and
# misclassified every French job/education entry as a date line. A year is a
# reliable, language-independent signal instead.
DATE_LINE_PATTERN = re.compile(r"(19|20)\d{2}")
ACCENT_COLOR = RGBColor(31, 78, 121)
ACCENT_HEX = "1F4E79"
PDF_ACCENT_COLOR = colors.HexColor("#1F4E79")


def _language_suffix(language: str) -> str:
    if language not in {"en", "fr"}:
        raise ValueError("language must be 'en' or 'fr'")
    return "_fr" if language == "fr" else ""


def _slugify(value: str, fallback: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", (value or "").lower()).strip("-")
    return slug or fallback


def _filename_component(value: str, fallback: str = "company") -> str:
    return re.sub(r"[^a-z0-9_-]+", "_", (value or "").lower()).strip("._-") or fallback


def _application_pack_directory(job: dict, output_dir: str | Path | None = None) -> Path:
    """Return a stable folder for one specific company, role, and posting URL."""
    target_dir = Path(output_dir) if output_dir is not None else OUTPUT_DIR / "resumes"
    company_slug = _slugify(str(job.get("company") or ""), "company")
    role_slug = _slugify(str(job.get("title") or ""), "role")
    url = str(job.get("url") or "").strip().lower()
    posting_id = sha256(url.encode("utf-8")).hexdigest()[:8] if url else "manual"
    return target_dir / f"{company_slug}__{role_slug}__{posting_id}"


@dataclass
class ResumeBlock:
    kind: str  # "header", "section", "summary", "entry", "date", "bullet"
    value: Any


@dataclass
class ParsedResume:
    header_name: str = ""
    contact_lines: list[str] = field(default_factory=list)
    blocks: list[ResumeBlock] = field(default_factory=list)
    raw_text: str = ""


class ResumeParser:
    """Parses plain text master resume into typed document blocks."""

    @staticmethod
    def parse(resume_text: str) -> ParsedResume:
        lines = [line.strip() for line in resume_text.splitlines()]
        first_section = next(
            (index for index, line in enumerate(lines) if line in SECTION_HEADERS), len(lines)
        )
        contact_lines_raw = [line for line in lines[1:first_section] if line]
        # Some master-resume files put phone/email/github each on its own line;
        # others combine them on one line separated by "|" (seen in the French
        # master resume). Normalize both into one flat list of individual items
        # so header rendering (and hyperlinking) works the same either way.
        contact_lines: list[str] = []
        for raw_line in contact_lines_raw:
            parts = [part.strip() for part in raw_line.split("|") if part.strip()]
            contact_lines.extend(parts or [raw_line])
        header_name = lines[0] if lines else ""

        blocks: list[ResumeBlock] = [ResumeBlock("header", (header_name, contact_lines))]
        index = first_section
        current_section = None

        while index < len(lines):
            line = lines[index]
            if not line:
                index += 1
                continue
            if line in SECTION_HEADERS:
                blocks.append(ResumeBlock("section", line.title()))
                current_section = line
                index += 1
                continue
            if line.startswith("-"):
                blocks.append(ResumeBlock("bullet", line[1:].strip()))
                index += 1
                continue
            if current_section not in SUMMARY_SECTION_HEADERS and DATE_LINE_PATTERN.search(line):
                blocks.append(ResumeBlock("date", line))
                index += 1
                continue
            # A plain line under Professional Summary or Languages is either
            # prose or a short fact ("French - Fluent"), not a job/education
            # heading, so it must not render bold the way an entry line does.
            if current_section in PLAIN_TEXT_SECTION_HEADERS:
                blocks.append(ResumeBlock("summary", line))
            else:
                blocks.append(ResumeBlock("entry", line))
            index += 1

        return ParsedResume(
            header_name=header_name,
            contact_lines=contact_lines,
            blocks=blocks,
            raw_text=resume_text,
        )


def _reorder_comma_list_by_relevance(items_line: str, job_description: str) -> str:
    """Reorder a comma-separated skills bullet so items the posting actually
    mentions come first. Never adds, removes, or rewords an item — it is a pure
    reordering of what's already true, purely to help a human/ATS skim see the
    overlap immediately."""
    items = [item.strip() for item in items_line.split(",") if item.strip()]
    jd_text = re.sub(r"[^a-z0-9+./#\s]", " ", (job_description or "").lower())
    if not items or not jd_text.strip():
        return items_line

    def is_relevant(item: str) -> bool:
        token = re.sub(r"[^a-z0-9+./#\s]", " ", item.lower()).strip()
        return bool(token) and token in jd_text

    matched = [item for item in items if is_relevant(item)]
    unmatched = [item for item in items if not is_relevant(item)]
    if not matched:
        return items_line
    return ", ".join(matched + unmatched)


def _apply_skill_reordering(resume_text: str, job_description: str) -> str:
    """Walk the raw resume text and reorder only the bullet lines that fall
    under a Technical Skills section, leaving every other line untouched."""
    if not job_description:
        return resume_text
    lines = resume_text.splitlines()
    output: list[str] = []
    current_section = None
    for line in lines:
        stripped = line.strip()
        if stripped in SECTION_HEADERS:
            current_section = stripped
            output.append(line)
            continue
        if current_section in SKILLS_SECTION_HEADERS and stripped.startswith("-"):
            bullet_text = stripped[1:].strip()
            reordered = _reorder_comma_list_by_relevance(bullet_text, job_description)
            output.append(f"- {reordered}")
            continue
        output.append(line)
    return "\n".join(output)


PROJECT_RELEVANCE_TERMS = {
    "plazio": (
        "cloud", "gcp", "google cloud", "azure", "terraform", "infrastructure as code",
        "iac", "pub/sub", "pubsub", "message queue", "devops", "testing", "test coverage",
        "qa", "automation", "ci/cd",
    ),
    "jobhunter": (
        "python", "automation", "scraping", "streamlit", "sqlite", "job search", "ats",
        "dashboard", "scheduler", "data pipeline", "fit score", "resume", "business analyst",
        "systems analyst", "requirements",
    ),
    "colorectal": (
        "machine learning", "deep learning", "pytorch", "torchvision", "computer vision",
        "image classification", "resnet", "scikit-learn", "healthcare", "medical",
    ),
    "472-a2": (
        "machine learning", "data science", "clustering", "decision tree", "mlp", "cnn",
        "fashion-mnist", "pytorch", "scikit-learn", "classification", "analytics", "pandas",
    ),
    "svelteship": (
        "svelte", "sveltekit", "javascript", "typescript", "full-stack", "web application",
        "logistics", "delivery", "prisma", "rest api", "relational database", "payment",
        "stripe", "business analyst", "product analyst",
    ),
    "campus guide": (
        "mobile", "react native", "expo", "navigation", "mapping", "maps", "directions",
        "calendar", "campus", "javascript", "typescript",
    ),
}


def _tailor_projects(resume_text: str, job_title: str, job_description: str) -> str:
    lines = resume_text.splitlines()
    project_header = next(
        (i for i, line in enumerate(lines) if line.strip() in {"PROJECTS", "PROJETS"}), None
    )
    if project_header is None:
        return resume_text
    section_end = next(
        (i for i in range(project_header + 1, len(lines)) if lines[i].strip() in SECTION_HEADERS),
        len(lines),
    )

    projects: list[list[str]] = []
    for line in lines[project_header + 1:section_end]:
        if line.strip() and not line.lstrip().startswith("-"):
            projects.append([line])
        elif projects:
            projects[-1].append(line)
    if not projects:
        return resume_text

    target = f"{job_title} {job_description}".lower()

    def score(project: list[str]) -> int:
        text = " ".join(project).lower()
        key = next((key for key in PROJECT_RELEVANCE_TERMS if key in text), None)
        if key is None:
            return 0
        return sum(term in target for term in PROJECT_RELEVANCE_TERMS[key])

    ranked = sorted(enumerate(projects), key=lambda pair: (-score(pair[1]), pair[0]))
    # Keep a second project when available so a sparse job description does not
    # produce a resume with an unnecessarily thin project section.
    selected = [project for _, project in ranked[:2]]
    return "\n".join(lines[:project_header + 1] + [line for project in selected for line in project] + lines[section_end:])


class IResumeExporter(ABC):
    """Abstract strategy for exporting resume documents in various file formats."""

    @abstractmethod
    def export(self, parsed: ParsedResume, output_path: Path) -> str:
        pass


class TxtResumeExporter(IResumeExporter):
    def export(self, parsed: ParsedResume, output_path: Path) -> str:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(parsed.raw_text, encoding="utf-8")
        return str(output_path)


def _add_docx_hyperlink(
    paragraph, url: str, text: str, font_size_pt: float = 9, color_hex: str = ACCENT_HEX
) -> None:
    """Insert a clickable hyperlink run into a python-docx paragraph.

    python-docx has no built-in hyperlink API, so this builds the required
    <w:hyperlink> XML directly. The run's font size/color are set here (not via
    paragraph.runs) because python-docx's `.runs` accessor only sees runs that
    are direct children of the paragraph, and a hyperlink run is nested one
    level deeper.
    """
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), color_hex)
    rpr.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)

    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(int(font_size_pt * 2)))  # half-points
    rpr.append(size)

    run.append(rpr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_contact_line_docx(document: Document, contacts: list[str], font_size_pt: float = 9) -> None:
    """Render the contact line, hyperlinking anything that looks like an email
    or a URL while leaving plain text (address, phone) as-is."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for index, item in enumerate(contacts):
        if index > 0:
            separator = paragraph.add_run("  |  ")
            separator.font.size = Pt(font_size_pt)
        if "@" in item:
            _add_docx_hyperlink(paragraph, f"mailto:{item}", item, font_size_pt=font_size_pt)
        elif item.lower().startswith(("http://", "https://")):
            _add_docx_hyperlink(paragraph, item, item, font_size_pt=font_size_pt)
        else:
            run = paragraph.add_run(item)
            run.font.size = Pt(font_size_pt)
    paragraph.paragraph_format.space_after = Pt(8)


class DocxResumeExporter(IResumeExporter):
    def export(self, parsed: ParsedResume, output_path: Path) -> str:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        _configure_docx(document)

        for block in parsed.blocks:
            kind, value = block.kind, block.value
            if kind == "header":
                name, contacts = value
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(name.upper())
                run.bold = True
                run.font.name = "Aptos Display"
                run.font.size = Pt(17)
                run.font.color.rgb = ACCENT_COLOR
                _add_contact_line_docx(document, contacts)
            elif kind == "section":
                paragraph = document.add_paragraph()
                run = paragraph.add_run(value.upper())
                run.bold = True
                run.font.size = Pt(10.5)
                run.font.color.rgb = ACCENT_COLOR
                _set_cell_border_bottom(paragraph)
                paragraph.paragraph_format.space_before = Pt(5)
                paragraph.paragraph_format.space_after = Pt(1)
            elif kind == "summary":
                # Regular justified prose -- not bold. Previously this fell
                # through to the "entry" branch below and rendered as a bold
                # heading-style line, which is wrong for a paragraph of text.
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                run = paragraph.add_run(value)
                run.font.size = Pt(9)
                paragraph.paragraph_format.space_after = Pt(4)
            elif kind == "entry":
                paragraph = document.add_paragraph()
                run = paragraph.add_run(value)
                run.bold = True
                run.font.size = Pt(9.5)
                paragraph.paragraph_format.space_before = Pt(3)
            elif kind == "date":
                paragraph = document.add_paragraph(value)
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.runs[0].italic = True
                paragraph.runs[0].font.size = Pt(8.5)
            elif kind == "bullet":
                paragraph = document.add_paragraph(style="Normal")
                paragraph.paragraph_format.left_indent = Inches(0.2)
                paragraph.paragraph_format.first_line_indent = Inches(-0.14)
                paragraph.paragraph_format.space_after = Pt(1)
                run = paragraph.add_run("•  " + value)
                run.font.size = Pt(9)

        document.save(output_path)
        return str(output_path)


class PdfResumeExporter(IResumeExporter):
    def export(self, parsed: ParsedResume, output_path: Path) -> str:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document = SimpleDocTemplate(
            str(output_path),
            pagesize=letter,
            topMargin=0.35 * inch,
            bottomMargin=0.35 * inch,
            leftMargin=0.65 * inch,
            rightMargin=0.65 * inch,
        )
        styles = _pdf_styles()
        story = []

        for block in parsed.blocks:
            kind, value = block.kind, block.value
            if kind == "header":
                name, contacts = value
                contact_markup = " | ".join(_pdf_contact_markup(item) for item in contacts)
                story.extend(
                    [
                        Paragraph(escape(name.upper()), styles["name"]),
                        Paragraph(contact_markup, styles["contact"]),
                        HRFlowable(
                            width="100%", thickness=0.9, color=PDF_ACCENT_COLOR, spaceAfter=2
                        ),
                    ]
                )
            elif kind == "section":
                story.extend(
                    [
                        Paragraph(escape(value.upper()), styles["section"]),
                        HRFlowable(
                            width="100%", thickness=0.5, color=PDF_ACCENT_COLOR, spaceAfter=1
                        ),
                    ]
                )
            elif kind == "summary":
                # Justified prose, not bold -- same fix as the DOCX exporter.
                story.append(Paragraph(escape(value), styles["summary"]))
            elif kind == "entry":
                story.append(Paragraph(escape(value), styles["entry"]))
            elif kind == "date":
                story.append(Paragraph(escape(value), styles["date"]))
            elif kind == "bullet":
                story.append(Paragraph("• " + escape(value), styles["bullet"]))

        document.build(story)
        return str(output_path)


def _pdf_contact_markup(item: str) -> str:
    """Return reportlab mini-markup for one contact item, hyperlinked if it's
    an email or URL, escaped either way so no stray text breaks the markup."""
    safe = escape(item)
    if "@" in item:
        return f'<link href="mailto:{safe}" color="#1F4E79"><u>{safe}</u></link>'
    if item.lower().startswith(("http://", "https://")):
        return f'<link href="{safe}" color="#1F4E79"><u>{safe}</u></link>'
    return safe


def build_master_resume(language: str = "en") -> str:
    """Load the English or French user-maintained master resume, falling back to
    a generic template ONLY if the real master file is missing.

    That fallback template is NOT reviewed per-application the way the master
    resume is -- if it ever gets used, warn loudly rather than silently sending
    out a generic/placeholder resume.
    """
    suffix = _language_suffix(language)
    resume_dir = Path(__file__).parent
    master_resume_path = resume_dir / f"master_resume{suffix}.txt"
    if master_resume_path.exists():
        return master_resume_path.read_text(encoding="utf-8")

    template_path = resume_dir / f"master_resume{suffix}.template.txt"
    if template_path.exists():
        warnings.warn(
            f"master_resume{suffix}.txt not found -- falling back to "
            f"master_resume{suffix}.template.txt. That template may contain "
            "placeholder/generic content (e.g. fake company names) and should "
            "NOT be sent to a real employer. Restore your real master resume "
            "file before generating any more application packs.",
            stacklevel=2,
        )
        return template_path.read_text(encoding="utf-8")

    return f"{USER_PROFILE.get('name', 'Your Name')}\n{USER_PROFILE.get('location', 'Your City, Region')}"


def tailor_resume_for_job(
    job_title: str, company_name: str, language: str = "en", job_description: str = ""
) -> str:
    """Emphasize verified skills and projects that best match a specific posting.

    The PROFESSIONAL SUMMARY is intentionally left as-authored: it's a general
    overview of who the candidate is, while PROJECTS (and skill ordering) are
    where per-posting tailoring happens.
    """
    resume_text = build_master_resume(language)
    target = f"{job_title}\n{job_description}".strip()
    if target:
        resume_text = _apply_skill_reordering(resume_text, target)
        resume_text = _tailor_projects(resume_text, job_title, job_description)
    return resume_text


def _resume_blocks(resume_text: str) -> list[tuple[str, object]]:
    """Convert the simple master-resume format into typed document blocks."""
    parsed = ResumeParser.parse(resume_text)
    return [(b.kind, b.value) for b in parsed.blocks]


def _set_cell_border_bottom(paragraph) -> None:
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_after = Pt(5)
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), ACCENT_HEX)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def _configure_docx(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(1.5)
    normal.paragraph_format.line_spacing = 1.0


def export_resume_docx(
    job_title: str,
    company_name: str,
    output_dir: str | Path | None = None,
    language: str = "en",
    job_description: str = "",
) -> str:
    """Export a clean, ATS-friendly resume based on Harvard's bullet template."""
    target_dir = Path(output_dir) if output_dir is not None else OUTPUT_DIR / "resumes"
    suffix = _language_suffix(language)
    company_slug = _filename_component(company_name)
    file_path = target_dir / f"resume_{company_slug}{suffix}.docx"
    resume_text = tailor_resume_for_job(job_title, company_name, language, job_description)
    parsed = ResumeParser.parse(resume_text)
    return DocxResumeExporter().export(parsed, file_path)


def _register_pdf_fonts() -> None:
    """Register a Unicode-capable font for PDF output on Windows."""
    fonts = {
        "ResumeArial": "C:\\Windows\\Fonts\\arial.ttf",
        "ResumeArial-Bold": "C:\\Windows\\Fonts\\arialbd.ttf",
        "ResumeArial-Italic": "C:\\Windows\\Fonts\\ariali.ttf",
        "ResumeArial-BoldItalic": "C:\\Windows\\Fonts\\arialbi.ttf",
    }
    for name, path in fonts.items():
        if name not in pdfmetrics.getRegisteredFontNames() and Path(path).exists():
            pdfmetrics.registerFont(TTFont(name, path))


def _pdf_styles() -> dict[str, ParagraphStyle]:
    _register_pdf_fonts()
    font_regular = "ResumeArial" if "ResumeArial" in pdfmetrics.getRegisteredFontNames() else "Helvetica"
    font_bold = (
        "ResumeArial-Bold"
        if "ResumeArial-Bold" in pdfmetrics.getRegisteredFontNames()
        else "Helvetica-Bold"
    )
    font_italic = (
        "ResumeArial-Italic"
        if "ResumeArial-Italic" in pdfmetrics.getRegisteredFontNames()
        else "Helvetica-Oblique"
    )
    base = getSampleStyleSheet()
    return {
        "name": ParagraphStyle(
            "resume-name",
            parent=base["Normal"],
            fontName=font_bold,
            fontSize=18,
            leading=21,
            textColor=PDF_ACCENT_COLOR,
            alignment=TA_CENTER,
            spaceAfter=2,
        ),
        "contact": ParagraphStyle(
            "resume-contact",
            parent=base["Normal"],
            fontName=font_regular,
            fontSize=8.5,
            leading=10,
            textColor=colors.HexColor("#404040"),
            alignment=TA_CENTER,
            spaceAfter=5,
        ),
        "section": ParagraphStyle(
            "resume-section",
            parent=base["Normal"],
            fontName=font_bold,
            fontSize=10.5,
            leading=12.5,
            textColor=PDF_ACCENT_COLOR,
            spaceBefore=6,
            spaceAfter=2,
        ),
        "summary": ParagraphStyle(
            "resume-summary",
            parent=base["Normal"],
            fontName=font_regular,
            fontSize=9,
            leading=11.5,
            alignment=TA_JUSTIFY,
            spaceAfter=4,
        ),
        "entry": ParagraphStyle(
            "resume-entry",
            parent=base["Normal"],
            fontName=font_bold,
            fontSize=9.8,
            leading=11.5,
            spaceBefore=3,
            spaceAfter=1,
        ),
        "date": ParagraphStyle(
            "resume-date",
            parent=base["Normal"],
            fontName=font_italic,
            fontSize=8.3,
            leading=9.5,
            textColor=colors.HexColor("#555555"),
            spaceAfter=0.5,
        ),
        "bullet": ParagraphStyle(
            "resume-bullet",
            parent=base["Normal"],
            fontName=font_regular,
            fontSize=9,
            leading=10.5,
            leftIndent=12,
            firstLineIndent=-8,
            spaceAfter=1,
        ),
        "body": ParagraphStyle(
            "resume-body",
            parent=base["Normal"],
            fontName=font_regular,
            fontSize=9,
            leading=10.5,
            spaceAfter=1.5,
        ),
    }


def export_resume_pdf(
    job_title: str,
    company_name: str,
    output_dir: str | Path | None = None,
    language: str = "en",
    job_description: str = "",
) -> str:
    """Export the same conservative, scan-friendly template as a PDF."""
    target_dir = Path(output_dir) if output_dir is not None else OUTPUT_DIR / "resumes"
    suffix = _language_suffix(language)
    company_slug = _filename_component(company_name)
    file_path = target_dir / f"resume_{company_slug}{suffix}.pdf"
    resume_text = tailor_resume_for_job(job_title, company_name, language, job_description)
    parsed = ResumeParser.parse(resume_text)
    return PdfResumeExporter().export(parsed, file_path)


def export_resume_files(
    job_title: str,
    company_name: str,
    formats=None,
    output_dir: str | Path | None = None,
    language: str = "en",
    job_description: str = "",
):
    target_dir = Path(output_dir) if output_dir is not None else OUTPUT_DIR / "resumes"
    target_dir.mkdir(parents=True, exist_ok=True)
    wanted_formats = formats or ["txt", "docx", "pdf"]
    result = {}
    suffix = _language_suffix(language)
    company_slug = _filename_component(company_name)
    resume_text = tailor_resume_for_job(job_title, company_name, language, job_description)
    parsed = ResumeParser.parse(resume_text)

    if "txt" in wanted_formats:
        txt_file = target_dir / f"resume_{company_slug}{suffix}.txt"
        result["txt"] = TxtResumeExporter().export(parsed, txt_file)
    if "docx" in wanted_formats:
        docx_file = target_dir / f"resume_{company_slug}{suffix}.docx"
        result["docx"] = DocxResumeExporter().export(parsed, docx_file)
    if "pdf" in wanted_formats:
        pdf_file = target_dir / f"resume_{company_slug}{suffix}.pdf"
        result["pdf"] = PdfResumeExporter().export(parsed, pdf_file)
    return result


def _job_requirement_summary(job_description: str) -> str:
    """Extract a short, presentation-safe set of role requirements."""
    plain_text = re.sub(r"<[^>]+>", " ", job_description or "")
    plain_text = re.sub(r"[*_`#]+", " ", plain_text)
    plain_text = re.sub(r"\s+", " ", plain_text).lower()
    terms = [
        ("python", "Python"),
        ("automation", "automation"),
        ("machine learning", "machine learning"),
        ("pytorch", "PyTorch"),
        ("scikit-learn", "scikit-learn"),
        ("pandas", "pandas"),
        ("numpy", "NumPy"),
        ("data pipeline", "data pipelines"),
        ("pytest", "pytest"),
        ("ci/cd", "CI/CD"),
        ("docker", "Docker"),
        ("sql", "SQL"),
        ("git", "Git"),
        ("apis", "APIs"),
        ("cloud", "cloud"),
        ("javascript", "JavaScript"),
        ("excel", "Excel"),
        ("business analysis", "business analysis"),
        ("requirements gathering", "requirements gathering"),
    ]
    matches = [(plain_text.index(term), label) for term, label in terms if term in plain_text]
    if not matches:
        return ""
    return ", ".join(label for _, label in sorted(matches)[:3])


def _verified_resume_summary(language: str) -> str:
    suffix = _language_suffix(language)
    resume_path = Path(__file__).parent / f"master_resume{suffix}.txt"
    if not resume_path.exists():
        return ""

    lines = resume_path.read_text(encoding="utf-8").splitlines()
    summary_headers = SUMMARY_SECTION_HEADERS
    in_summary = False
    summary_lines = []
    for line in lines:
        stripped = line.strip()
        if stripped in summary_headers:
            in_summary = True
            continue
        if in_summary and stripped in SECTION_HEADERS:
            break
        if in_summary and stripped:
            summary_lines.append(stripped)
    return " ".join(summary_lines)


def generate_cover_letter(job_title: str, company_name: str, job_description: str = "") -> str:
    profile = USER_PROFILE["name"]
    location = USER_PROFILE["location"]
    title = job_title or "Software Engineer"
    company = company_name or "Hiring Company"
    requirements = _job_requirement_summary(job_description)
    resume_summary = _verified_resume_summary("en")
    evidence = (
        f"My professional summary is: {resume_summary} "
        if resume_summary
        else "I have included my resume for your review. "
    )
    description_note = (
        f" The posting highlights {requirements}; I would welcome a discussion about the role's priorities."
        if requirements
        else ""
    )
    return (
        f"Dear Hiring Manager at {company},\n\n"
        f"I am writing to express my interest in the {title} opportunity at {company}. "
        f"{evidence}{description_note}\n\n"
        "I look forward to the opportunity to discuss how my background aligns with your needs.\n\n"
        f"Sincerely,\n{profile}\n{location}"
    )


def generate_cover_letter_french(
    job_title: str, company_name: str, job_description: str = ""
) -> str:
    """Generate a concise French cover letter using the user's verified background."""
    title = job_title or "poste recherché"
    company = company_name or "votre organisation"
    requirements = _job_requirement_summary(job_description)
    resume_summary = _verified_resume_summary("fr")
    evidence = (
        f"Mon profil professionnel est le suivant : {resume_summary} "
        if resume_summary
        else "Vous trouverez mon curriculum vitæ en pièce jointe. "
    )
    description_note = (
        f" L'offre met l'accent sur {requirements}; je serais heureux d'échanger sur les priorités du poste."
        if requirements
        else ""
    )
    return (
        f"Madame, Monsieur,\n\n"
        f"Je vous soumets ma candidature au poste de {title} chez {company}. "
        f"{evidence}{description_note}\n\n"
        "Je serais heureux de discuter de la façon dont mon parcours répond aux besoins de ce poste.\n\n"
        f"Cordialement,\n{USER_PROFILE['name']}\n{USER_PROFILE['location']}"
    )


def export_cover_letter(
    job_title: str,
    company_name: str,
    job_description: str = "",
    output_dir: str | Path | None = None,
    language: str = "en",
) -> str:
    target_dir = Path(output_dir) if output_dir is not None else OUTPUT_DIR / "resumes"
    target_dir.mkdir(parents=True, exist_ok=True)
    suffix = _language_suffix(language)
    company_slug = _filename_component(company_name)
    file_path = target_dir / f"cover_letter_{company_slug}{suffix}.docx"
    document = Document()
    _configure_docx(document)

    letter = (
        generate_cover_letter_french(job_title, company_name, job_description)
        if language == "fr"
        else generate_cover_letter(job_title, company_name, job_description)
    )
    for text in letter.split("\n\n"):
        paragraph = document.add_paragraph(text)
        paragraph.paragraph_format.space_after = Pt(9)
        paragraph.paragraph_format.line_spacing = 1.08

    document.save(file_path)
    return str(file_path)


def generate_application_pack(
    job: dict,
    output_dir: str | Path | None = None,
) -> dict[str, str]:
    """
    Explicitly generate tailored resume (PDF, DOCX) and cover letter (DOCX) files
    for both English and French for a specific job when requested.
    """
    job_title = job.get("title", "Software Engineer")
    company_name = job.get("company", "Company")
    description = job.get("description", "")
    pack_dir = _application_pack_directory(job, output_dir)

    # Export English resume & cover letter
    en_files = export_resume_files(
        job_title,
        company_name,
        formats=["docx", "pdf"],
        output_dir=pack_dir,
        language="en",
        job_description=description,
    )
    en_cover = export_cover_letter(
        job_title, company_name, job_description=description, output_dir=pack_dir, language="en"
    )

    # Export French resume & cover letter
    fr_files = export_resume_files(
        job_title,
        company_name,
        formats=["docx", "pdf"],
        output_dir=pack_dir,
        language="fr",
        job_description=description,
    )
    fr_cover = export_cover_letter(
        job_title, company_name, job_description=description, output_dir=pack_dir, language="fr"
    )

    return {
        "en_resume_docx": en_files.get("docx", ""),
        "en_resume_pdf": en_files.get("pdf", ""),
        "en_cover_letter_docx": en_cover,
        "fr_resume_docx": fr_files.get("docx", ""),
        "fr_resume_pdf": fr_files.get("pdf", ""),
        "fr_cover_letter_docx": fr_cover,
    }


def get_application_pack_files(
    job: dict,
    output_dir: str | Path | None = None,
) -> dict[str, str | bool]:
    """
    Inspect which resume and cover letter files exist on disk for a job. Does NOT
    generate files. `exists` means at least one file is available; `complete`
    means every expected file was generated.
    """
    target_dir = _application_pack_directory(job, output_dir)
    company_slug = _filename_component(str(job.get("company") or ""))

    en_pdf = target_dir / f"resume_{company_slug}.pdf"
    en_docx = target_dir / f"resume_{company_slug}.docx"
    en_cover = target_dir / f"cover_letter_{company_slug}.docx"
    fr_pdf = target_dir / f"resume_{company_slug}_fr.pdf"
    fr_docx = target_dir / f"resume_{company_slug}_fr.docx"
    fr_cover = target_dir / f"cover_letter_{company_slug}_fr.docx"

    expected_files = (en_pdf, en_docx, en_cover, fr_pdf, fr_docx, fr_cover)
    complete = all(path.exists() for path in expected_files)
    exists = any(path.exists() for path in expected_files)

    return {
        "exists": exists,
        "complete": complete,
        "en_resume_docx": str(en_docx) if en_docx.exists() else "",
        "en_resume_pdf": str(en_pdf) if en_pdf.exists() else "",
        "en_cover_letter_docx": str(en_cover) if en_cover.exists() else "",
        "fr_resume_docx": str(fr_docx) if fr_docx.exists() else "",
        "fr_resume_pdf": str(fr_pdf) if fr_pdf.exists() else "",
        "fr_cover_letter_docx": str(fr_cover) if fr_cover.exists() else "",
    }


_ATS_STOP_WORDS = {
    "about", "after", "also", "and", "are", "as", "at", "be", "been", "but", "by",
    "can", "candidate", "company", "for", "from", "have", "has", "in", "into", "is",
    "it", "its", "job", "of", "on", "or", "our", "role", "that", "the", "their",
    "team", "this", "to", "we", "with", "you", "your",
}


def _ats_terms(text: str) -> set[str]:
    plain_text = re.sub(r"<[^>]+>", " ", text or "").lower()
    return {
        token.strip("./-")
        for token in re.findall(r"[a-z][a-z0-9+#./-]*", plain_text)
        if len(token.strip("./-")) > 1 and token.strip("./-") not in _ATS_STOP_WORDS
    }


def ats_analysis(resume_text: str, job_description: str) -> dict:
    """Estimate informative word-token coverage; this is not an ATS prediction."""
    resume_keywords = _ats_terms(resume_text)
    jd_tokens = _ats_terms(job_description)
    missing_keywords = sorted(jd_tokens - resume_keywords)
    description = (job_description or "").lower()
    missing_skills = [
        keyword
        for keyword in USER_PROFILE["keywords"]
        if keyword.lower() in description and keyword.lower() not in (resume_text or "").lower()
    ]
    matched_count = len(jd_tokens & resume_keywords)
    ats_score = round(100 * matched_count / len(jd_tokens)) if jd_tokens else 0
    return {
        "ats_score": ats_score,
        "score_type": "keyword coverage estimate",
        "missing_keywords": missing_keywords[:15],
        "missing_skills": missing_skills[:10],
        "improvements": [
            "Add role-specific keywords from the job description.",
            "Highlight your Python, automation, and machine learning project experience.",
            "Emphasize PyTorch, scikit-learn, SQL, and automation scripting skills where relevant.",
        ],
    }
